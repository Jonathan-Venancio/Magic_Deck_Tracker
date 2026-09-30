from dataclasses import dataclass
from io import BytesIO

from fastapi import HTTPException
from odf.opendocument import OpenDocumentSpreadsheet, load
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import Workbook, load_workbook

from app.mana import cell_text, fold_text, header_key

HEADERS = [
    "Nome da carta",
    "Coleção",
    "Número da carta",
    "Quantidade",
    "Tipo",
    "Custo de mana",
    "Categoria",
    "Cores",
    "Descrição / Tradução",
]

HEADER_FIELDS = {
    "nome": "name",
    "nome da carta": "name",
    "name": "name",
    "colecao": "collection",
    "collection": "collection",
    "numero": "number",
    "numero da carta": "number",
    "number": "number",
    "quantidade": "quantity",
    "qtd": "quantity",
    "qty": "quantity",
    "quantity": "quantity",
    "tipo": "type_line",
    "type": "type_line",
    "custo de mana": "mana",
    "custo": "mana",
    "mana": "mana",
    "mana cost": "mana",
    "categoria": "category",
    "category": "category",
    "cores": "colors",
    "cor": "colors",
    "colors": "colors",
    "color": "colors",
    "descricao": "text",
    "traducao": "text",
    "descricao traducao": "text",
    "texto": "text",
    "text": "text",
}

SAMPLE_ROWS = [
    [
        "Katara, Mestra da Água",
        "Avatar",
        "127",
        "1",
        "Criatura Lendária — Humano Mago",
        "2U",
        "Criatura",
        "U",
        "Dois genéricos + um azul. Apague as linhas de exemplo e coloque as suas cartas.",
    ],
    [
        "Sol Ring",
        "Commander",
        "1",
        "1",
        "Artefato",
        "1",
        "Outro",
        "",
        "Um genérico (mana de qualquer cor).",
    ],
    [
        "Geist of Saint Traft",
        "Innistrad",
        "213",
        "1",
        "Criatura Lendária — Espírito Clérigo",
        "1WU",
        "Criatura",
        "W, U",
        "Um genérico + branco + azul. Duas cores.",
    ],
    [
        "Ilha",
        "Avatar",
        "287",
        "8",
        "Terreno Básico — Ilha",
        "",
        "Terreno",
        "",
        "Terreno não tem custo de mana.",
    ],
]

INSTRUCTIONS = [
    "Baixe este modelo, preencha a aba Cartas e envie de volta no app.",
    "A imagem da carta não entra na planilha: você adiciona depois, carta a carta.",
    "",
    "Coleção: se o nome ainda não existir, o app cria. Se já existir, as cartas entram nela.",
    "Número da carta: o número impresso. Não pode repetir na mesma coleção; se repetir, a carta é atualizada (a foto permanece).",
    "",
    "Custo de mana — junte o número e as letras, sem espaço:",
    "  O número é mana genérica (qualquer cor). 1 = um genérico. 2 = dois genéricos.",
    "  As letras são mana colorida: W branco, U azul, B preto, R vermelho, G verde.",
    "  2U  = dois genéricos + um azul",
    "  1   = um genérico",
    "  1WU = um genérico + branco + azul (carta de duas cores)",
    "  UU  = dois azuis",
    "  R   = um vermelho",
    "  (vazio) = terreno, sem custo",
    "",
    "Cores: pode deixar vazio — o app lê as cores do custo (2U vira Azul).",
    "  Várias cores: U, W   ou   Azul, Branco   ou   UW",
    "",
    "Categoria: Criatura, Mágica, Terreno ou Outro.",
    "",
    "Excel: salve como .xlsx. LibreOffice Calc: .ods (não use .odt, que é arquivo de texto).",
]

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 2000


@dataclass
class SheetRow:
    row: int
    name: str
    collection: str
    number: str
    quantity: str
    type_line: str
    mana: str
    category: str
    colors: str
    text: str


def _is_instructions(name: str) -> bool:
    return fold_text(name).startswith("instru")


def build_template(fmt: str) -> tuple[bytes, str, str]:
    if fmt == "xlsx":
        return _build_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "modelo-cartas.xlsx"
    if fmt == "ods":
        return _build_ods(), "application/vnd.oasis.opendocument.spreadsheet", "modelo-cartas.ods"
    raise HTTPException(status_code=400, detail="Formato do modelo inválido.")


def read_card_rows(filename: str, content: bytes) -> list[SheetRow]:
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="Arquivo grande demais. Use no máximo 8 MB.")
    name = (filename or "").rsplit("/", 1)[-1].lower()
    if name.endswith(".odt"):
        raise HTTPException(
            status_code=400,
            detail="Arquivo .odt é do Writer. No LibreOffice Calc, salve como .ods ou .xlsx.",
        )
    if name.endswith(".xlsx"):
        tables = _read_xlsx(content)
    elif name.endswith(".ods"):
        tables = _read_ods(content)
    else:
        raise HTTPException(status_code=400, detail="Envie um Excel (.xlsx) ou uma planilha do Calc (.ods).")
    rows = _rows_from_tables(tables)
    if not rows:
        raise HTTPException(status_code=400, detail="A planilha não tem cartas para importar.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"A planilha tem cartas demais. Máximo: {MAX_ROWS}.")
    return rows


def _build_xlsx() -> bytes:
    workbook = Workbook()
    cards = workbook.active
    cards.title = "Cartas"
    cards.append(HEADERS)
    for row in SAMPLE_ROWS:
        cards.append(row)
    for index, header in enumerate(HEADERS, start=1):
        cards.column_dimensions[chr(64 + index)].width = min(28, max(14, len(header) + 4))
    info = workbook.create_sheet("Instruções")
    info.column_dimensions["A"].width = 92
    for line in INSTRUCTIONS:
        info.append([line])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _add_ods_cell(row: TableRow, text: str) -> None:
    cell = TableCell(valuetype="string")
    cell.addElement(P(text=text or " "))
    row.addElement(cell)


def _build_ods() -> bytes:
    document = OpenDocumentSpreadsheet()
    cards = Table(name="Cartas")
    header_row = TableRow()
    for header in HEADERS:
        _add_ods_cell(header_row, header)
    cards.addElement(header_row)
    for values in SAMPLE_ROWS:
        row = TableRow()
        for value in values:
            _add_ods_cell(row, value)
        cards.addElement(row)
    info = Table(name="Instruções")
    for line in INSTRUCTIONS:
        row = TableRow()
        _add_ods_cell(row, line)
        info.addElement(row)
    document.spreadsheet.addElement(cards)
    document.spreadsheet.addElement(info)
    buffer = BytesIO()
    document.write(buffer)
    return buffer.getvalue()


def _read_xlsx(content: bytes) -> dict[str, list[list[str]]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise HTTPException(status_code=400, detail="Não foi possível ler o Excel.") from error
    tables: dict[str, list[list[str]]] = {}
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([cell_text(cell) for cell in row])
        tables[sheet.title] = rows
    return tables


def _ods_cell_text(cell: TableCell) -> str:
    parts = [str(paragraph) for paragraph in cell.getElementsByType(P)]
    text = "\n".join(part.strip() for part in parts if part and part.strip())
    if text:
        return text
    raw = cell.getAttribute("value")
    if not raw:
        return ""
    try:
        return cell_text(float(raw))
    except ValueError:
        return str(raw).strip()


def _ods_row_values(row: TableRow) -> list[str]:
    values: list[str] = []
    for cell in row.getElementsByType(TableCell):
        repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        text = _ods_cell_text(cell)
        values.extend([text] * min(max(repeat, 1), 20))
        if len(values) >= 20:
            break
    return values[:20]


def _read_ods(content: bytes) -> dict[str, list[list[str]]]:
    try:
        document = load(BytesIO(content))
    except Exception as error:
        raise HTTPException(status_code=400, detail="Não foi possível ler a planilha do Calc.") from error
    tables: dict[str, list[list[str]]] = {}
    for table in document.spreadsheet.getElementsByType(Table):
        name = str(table.getAttribute("name") or "Planilha")
        rows = [_ods_row_values(row) for row in table.getElementsByType(TableRow)]
        tables[name] = rows
    return tables


def _pick_table(tables: dict[str, list[list[str]]]) -> list[list[str]]:
    for name, rows in tables.items():
        if fold_text(name) == "cartas":
            return rows
    for name, rows in tables.items():
        if not _is_instructions(name):
            return rows
    raise HTTPException(status_code=400, detail="Não achei a aba de cartas na planilha.")


def _map_headers(header_row: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(header_row):
        field = HEADER_FIELDS.get(header_key(header))
        if field and field not in mapping:
            mapping[field] = index
    missing = [label for label, field in (("Nome da carta", "name"), ("Coleção", "collection"), ("Número da carta", "number")) if field not in mapping]
    if missing:
        raise HTTPException(status_code=400, detail=f"Faltam colunas: {', '.join(missing)}.")
    return mapping


def _rows_from_tables(tables: dict[str, list[list[str]]]) -> list[SheetRow]:
    grid = _pick_table(tables)
    if not grid:
        return []
    mapping = _map_headers(grid[0])

    def value(cells: list[str], field: str) -> str:
        index = mapping.get(field)
        if index is None or index >= len(cells):
            return ""
        return cells[index]

    rows: list[SheetRow] = []
    for offset, cells in enumerate(grid[1:], start=2):
        if not any(cells):
            continue
        rows.append(
            SheetRow(
                row=offset,
                name=value(cells, "name"),
                collection=value(cells, "collection"),
                number=value(cells, "number"),
                quantity=value(cells, "quantity"),
                type_line=value(cells, "type_line"),
                mana=value(cells, "mana"),
                category=value(cells, "category"),
                colors=value(cells, "colors"),
                text=value(cells, "text"),
            )
        )
    return rows
